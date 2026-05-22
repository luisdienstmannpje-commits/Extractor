import hashlib
import json
import os
import re
from datetime import datetime, date, timedelta
from typing import Any, Optional
from services.sentence_finder import extract_sentence_from_pdf
from services.ai_client import extract_data_with_gemini
from services.pre_extractor import pre_extract
from services.text_processor import find_section_hybrid
from services.database import (
    get_user_credits,
    deduct_credit,
    get_cache_repo,
    get_extraction_repo,
    quota_excedida,
)
from services.legal_validator import validar_dados_completo
from services.legal_engine.rule_registry import carregar_todas_as_regras   # Gap 2
from services.legal_engine.engine import LegalRuleEngine                   # Gap 2
from services.legal_engine.dynamic_rule_loader import (                    # Self-Healing
    carregar_regras_ativas,
    executar_shadow_pipeline,
)
from services.explanation_engine import (
    gerar_explicacoes,
    gerar_parecer_parcelas_apuradas,
    obter_textos_padrao_criterios_parecer,
    gerar_parecer_tecnico_completo,
)
from services.jurisprudencia.consistencia.verba_deduplicator import deduplicar_verbas  # S11
from memoria_calculo import gerar_memoria                                               # M1
from models import ProcessoTrabalhista
from config import settings
from services.extraction_engine import enriquecer_para_raiox
from services.verba_page_anchor import anchor_verbas_to_pages, build_page_chunks
from services.cache_key import pdf_cache_storage_key
from services.memorial_pedidos import gerar_memorial_pedidos, gerar_memorial_defesa

# ── Singleton — carregado uma vez na inicialização do módulo ─────────────────
# Evita recarregar as 20 regras e reordenar por prioridade a cada request.
_RULE_ENGINE = LegalRuleEngine(carregar_todas_as_regras())

# ── Caminho da pasta de playbooks ────────────────────────────────────────────
SKILLS_DIR = os.path.join(os.path.dirname(__file__), "..", "skills")

def _load_skill(filename: str) -> str:
    """Carrega um playbook .md da pasta skills/."""
    path = os.path.join(SKILLS_DIR, filename)
    if os.path.exists(path):
        with open(path, "r", encoding="utf-8") as f:
            return f.read()
    print(f"[SKILL] Aviso: playbook não encontrado: {filename}", flush=True)
    return ""

def _hash_pdf(file_bytes: bytes) -> str:
    return hashlib.sha256(file_bytes).hexdigest()


def _hash_dossie(files_list: list) -> str:
    """Hash composto para cache de dossiê: ordenado por nome para determinismo."""
    h = hashlib.sha256()
    for (name, content) in sorted(files_list, key=lambda x: (x[0],)):
        h.update(name.encode("utf-8", errors="replace"))
        h.update(content)
    return h.hexdigest()


def _extrair_texto_arquivo_dossie(filename: str, file_bytes: bytes) -> str:
    """
    Extrai texto de um arquivo para o super-contexto do dossiê.
    PDF → sentence_finder; DOCX/DOC → python-docx ou learning_engine; XLSX/XLS → openpyxl/pandas;
    PJC/XML → decode; JPG/PNG → Gemini multimodal (OCR/descrição).
    """
    import io
    fn = (filename or "").lower()
    if fn.endswith(".pdf"):
        try:
            texto, _ = extract_sentence_from_pdf(file_bytes)
            return texto or ""
        except Exception as e:
            print(f"[PROCESSOR] Erro ao extrair PDF {filename}: {e}", flush=True)
            return ""
    if fn.endswith(".docx"):
        try:
            from docx import Document
            doc = Document(io.BytesIO(file_bytes))
            paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            if paras:
                return "\n".join(paras)
        except Exception:
            pass
        try:
            from services.learning_engine import _extrair_texto_docx
            return _extrair_texto_docx(file_bytes) or ""
        except Exception as e:
            print(f"[PROCESSOR] Erro ao extrair DOCX {filename}: {e}", flush=True)
        return ""
    if fn.endswith(".doc"):
        try:
            raw = file_bytes.decode("latin-1", errors="ignore")
            linhas = [l.strip() for l in raw.split("\n") if len(l.strip()) > 20]
            return "\n".join(linhas[:200])
        except Exception as e:
            print(f"[PROCESSOR] Erro ao extrair DOC {filename}: {e}", flush=True)
        return ""
    if fn.endswith(".xlsx") or fn.endswith(".xls"):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
            linhas = []
            for sheet in wb.worksheets:
                linhas.append(f"=== Planilha: {sheet.title} ===")
                for row in sheet.iter_rows(values_only=True):
                    celulas = [str(c).strip() if c is not None else "" for c in row]
                    if any(celulas):
                        linhas.append("\t".join(celulas))
            wb.close()
            return "\n".join(linhas) if linhas else ""
        except Exception as e:
            print(f"[PROCESSOR] Erro ao extrair Excel {filename}: {e}", flush=True)
            return f"[Planilha {filename}: não foi possível ler. Use .xlsx.]"
    if fn.endswith(".pjc") or fn.endswith(".xml"):
        try:
            text = file_bytes.decode("iso-8859-1", errors="replace")
            if len(text) > 120000:
                text = text[:120000] + "\n[... XML truncado ...]"
            return text
        except Exception:
            try:
                return file_bytes.decode("utf-8", errors="replace")[:120000]
            except Exception:
                return ""
    if fn.endswith(".jpg") or fn.endswith(".jpeg") or fn.endswith(".png"):
        try:
            from services.ai_client import extrair_texto_ou_descricao_imagem
            mime = "image/png" if fn.endswith(".png") else "image/jpeg"
            return extrair_texto_ou_descricao_imagem(file_bytes, mime_type=mime)
        except Exception as e:
            print(f"[PROCESSOR] Erro ao processar imagem {filename}: {e}", flush=True)
            return f"[Imagem {filename}: erro ao extrair texto]"
    return ""

# ---------------------------------------------------------------------------
# Campos de texto simples do ProcessoTrabalhista
# ---------------------------------------------------------------------------
CAMPOS_TEXTO = [
    # Identificação
    "numero_processo",
    "vara_trabalho",
    "reclamante",
    "reclamada",
    "tipo_rito",
    "funcao_reclamante",
    "data_sentenca",
    "data_ajuizamento",
    "advogado_reclamante",
    "advogado_reclamada",
    "juiz_responsavel",
    "valor_causa",
    # Contrato
    "data_admissao",
    "data_demissao",
    "motivo_rescisao",
    "tipo_contrato",
    "salario_base",
    "jornada_contratual",
    "horario_trabalho",
    "aviso_previo_dias",
    "data_saida_ctps",
    "anotacao_ctps",
    "seguro_desemprego",
    # Parâmetros de cálculo
    "indice_correcao",
    "juros_mora",
    "contribuicao_previdenciaria",
    "ir_retido_fonte",
    "honorarios_sucumbenciais",
    "percentual_honorarios",
    "custas_processuais",
    # Valores fixos e multas
    "dano_moral",
    "dano_material",
    "multa_art_467",
    "multa_art_477",
    # FGTS especial
    "fgts_sobre_aviso_previo",
    "fgts_multa_40_aviso_previo",
    "fgts_sobre_ferias_indenizadas",
    "fgts_periodo_completo",
    "fgts_observacoes",
]

# Campos de texto simples de cada VerbaDeferida
CAMPOS_VERBA_TEXTO = [
    "nome",
    "status_final",
    "periodo",
    "percentual",
    "quantidade_diaria",
    "base_calculo",
    "valor_fixado",
    "observacoes",
]

# Valores que a IA costuma retornar no lugar de null
SUSPICIOUS = frozenset([
    "não informado", "nao informado", "n/a", "null", "none",
    "não consta", "nao consta", "desconhecido", "indefinido",
    "não identificado", "nao identificado", "-", "",
])

def _clean_str(value) -> str | None:
    """
    Retorna None se o valor for suspeito/vazio; caso contrário, retorna uma string.

    Importante: a IA às vezes retorna números (float/int) para campos textuais
    como `quantidade_diaria` ou `percentual`. Para não estourar validação Pydantic,
    convertemos qualquer valor não-nulo em string, mantendo apenas os marcadores
    "não informado"/etc. como None.
    """
    if value is None:
        return None
    if isinstance(value, str):
        v = value.strip()
        if v.lower() in SUSPICIOUS:
            return None
        return v
    # Qualquer outro tipo (int, float, etc.) é convertido para string
    return str(value)


def _dedup_alertas(*listas: list[str]) -> list[str]:
    """
    Mescla listas de alertas removendo duplicados, preservando a ordem
    de primeira aparição. Útil para evitar que o mesmo alerta jurídico
    apareça duas vezes quando proveniente de mais de uma camada.
    """
    vistos: set[str] = set()
    resultado: list[str] = []
    for lista in listas:
        if not lista:
            continue
        for alerta in lista:
            if alerta not in vistos:
                vistos.add(alerta)
                resultado.append(alerta)
    return resultado

def _clean_bool(value, default: bool) -> bool:
    """Garante que o valor seja booleano. Aceita variações semânticas (ex: justiça gratuita 'concedida' -> True)."""
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        v = value.lower().strip()
        if v in ("true", "sim", "yes", "1", "concedida", "concedido", "deferida", "deferido"):
            return True
        if v in ("false", "não", "nao", "no", "0", "indeferida", "indeferido", "negada", "negado"):
            return False
    return default


# Chaves HIGH emitidas pelo PreExtractor — só estas podem sobrescrever a saída da IA
_HIGH_MERGE_KEYS = frozenset(
    {
        "numero_processo",
        "reclamante",
        "reclamada",
        "vara_trabalho",
        "data_sentenca",
        "justica_gratuita",
        "tipo_rito",
    }
)


def _aplicar_pre_high(cleaned: dict, pre_high: dict | None) -> dict:
    """
    Sobrescreve campos já validados com valores do pré-extrator (alta confiança).

    Regras: ver `test_extraction_cycle_merge_pre_high.py` e AI_AGENT_EXTRACTION_CYCLE.md.
    Emite uma linha `[PRE-HIGH]` em stdout quando algum campo whitelist muda de fato.
    """
    if not pre_high:
        return cleaned
    alterados: list[str] = []
    for k, v in pre_high.items():
        if k not in _HIGH_MERGE_KEYS:
            continue
        if v is None:
            continue
        if isinstance(v, str) and not v.strip():
            continue
        if k == "justica_gratuita":
            novo = bool(v)
        else:
            novo = v.strip() if isinstance(v, str) else v
        antigo = cleaned.get(k)
        if novo != antigo:
            alterados.append(k)
        cleaned[k] = novo
    if alterados:
        print("[PRE-HIGH] " + ", ".join(alterados), flush=True)
    return cleaned


def _validate_result(data: dict) -> dict:
    """
    Validação pós-IA:
    - Limpa strings suspeitas → None
    - Garante booleanos corretos
    - Garante listas onde esperado
    - Normaliza campos de cada verba
    """
    cleaned = {}

    # ── Campos de texto simples ───────────────────────────────────────────────
    for campo in CAMPOS_TEXTO:
        cleaned[campo] = _clean_str(data.get(campo))

    # ── Booleanos ─────────────────────────────────────────────────────────────
    cleaned["justica_gratuita"] = _clean_bool(data.get("justica_gratuita"), default=False)

    # ── Lista de verbas ───────────────────────────────────────────────────────
    verbas_raw = data.get("verbas_deferidas", [])
    if not isinstance(verbas_raw, list):
        verbas_raw = []

    verbas_limpas = []
    for verba in verbas_raw:
        if not isinstance(verba, dict):
            continue

        v = {}

        # Campos de texto
        for campo in CAMPOS_VERBA_TEXTO:
            v[campo] = _clean_str(verba.get(campo))

        # Nome nunca pode ser null
        if not v.get("nome"):
            v["nome"] = "Verba não identificada"

        # status_final tem valor padrão
        if not v.get("status_final"):
            v["status_final"] = "não informado"

        # integracao_salarial: booleano ou null
        integ = verba.get("integracao_salarial")
        if isinstance(integ, bool):
            v["integracao_salarial"] = integ
        elif isinstance(integ, str) and integ.lower() in ("true", "sim"):
            v["integracao_salarial"] = True
        elif isinstance(integ, str) and integ.lower() in ("false", "não", "nao"):
            v["integracao_salarial"] = False
        else:
            v["integracao_salarial"] = None

        # reflexos: lista de strings limpas
        reflexos_raw = verba.get("reflexos", [])
        if isinstance(reflexos_raw, list):
            v["reflexos"] = [r for r in reflexos_raw if isinstance(r, str) and r.strip()]
        else:
            v["reflexos"] = []

        # P3: frações de avos (ex: "11/12", "2/12") pertencem ao campo período,
        # não à quantidade diária. A IA frequentemente confunde os dois.
        qtd = str(v.get("quantidade_diaria") or "")
        if re.search(r"^\d+/12$", qtd.strip()):
            if not v.get("periodo"):
                v["periodo"] = qtd.strip()
            v["quantidade_diaria"] = None

        verbas_limpas.append(v)

    cleaned["verbas_deferidas"] = verbas_limpas

    # ── Pós-processamento determinístico (zero tokens) ────────────────────────

    # 1. jornada_contratual — deriva do horario_trabalho quando a IA não extraiu
    if not cleaned.get("jornada_contratual") and cleaned.get("horario_trabalho"):
        h = cleaned["horario_trabalho"]
        m = re.search(
            r"(\d{1,2})(?::(\d{2}))?(?:h\d*)?\s*[àa][s]?\s*(\d{1,2})(?::(\d{2}))?(?:h\d*)?"
            r".*?com\s+(\d+(?:[,\.]\d+)?)\s*(?:h(?:ora)?|:00)",
            h, re.IGNORECASE
        )
        if m:
            try:
                entrada_h = int(m.group(1)); entrada_m = int(m.group(2) or 0)
                saida_h   = int(m.group(3)); saida_m   = int(m.group(4) or 0)
                intervalo = float(m.group(5).replace(",", "."))
                total_min = (saida_h * 60 + saida_m) - (entrada_h * 60 + entrada_m) - intervalo * 60
                diarias   = total_min / 60
                semanais  = diarias * 6
                if 0 < diarias <= 12:
                    cleaned["jornada_contratual"] = f"{diarias:.0f}h diárias / {semanais:.0f}h semanais"
                    print(f"[POSTP] jornada_contratual derivada: {cleaned['jornada_contratual']}", flush=True)
            except Exception:
                pass

    # 2. fgts_periodo_completo — completa com datas reais quando genérico
    fgts = cleaned.get("fgts_periodo_completo") or ""
    if fgts and "/" not in fgts:
        adm   = cleaned.get("data_admissao")
        saida = cleaned.get("data_saida_ctps") or cleaned.get("data_demissao")
        if adm and saida:
            cleaned["fgts_periodo_completo"] = f"Todo o período contratual — {adm} a {saida}"
            print(f"[POSTP] fgts_periodo_completo completado: {cleaned['fgts_periodo_completo']}", flush=True)

    # 3. prescricao_quinquenal — calcula automaticamente (ajuizamento - 5 anos)
    if not cleaned.get("prescricao_quinquenal") and cleaned.get("data_ajuizamento"):
        try:
            d, m_n, y = cleaned["data_ajuizamento"].split("/")
            ajuiz = date(int(y), int(m_n), int(d))
            prescricao = ajuiz.replace(year=ajuiz.year - 5)
            cleaned["prescricao_quinquenal"] = prescricao.strftime("%d/%m/%Y")
            print(f"[POSTP] prescricao_quinquenal calculada: {cleaned['prescricao_quinquenal']}", flush=True)
        except Exception:
            pass

    # 4. divisor_horas — detecta por regex no texto do dispositivo / jornada
    if not cleaned.get("divisor_horas"):
        textos_busca = [
            cleaned.get("horario_trabalho") or "",
            cleaned.get("jornada_contratual") or "",
        ]
        for v in cleaned.get("verbas_deferidas", []):
            nome = (v.get("nome") or "").lower()
            if "hora" in nome and "extra" in nome:
                textos_busca.append(v.get("base_calculo") or "")
                textos_busca.append(v.get("observacoes") or "")

        texto_concat = " ".join(textos_busca)
        m_div = re.search(r"\b(150|180|200|220)\s*(?:h(?:oras?)?|\/\s*m[eê]s)?\b", texto_concat, re.IGNORECASE)
        if m_div:
            cleaned["divisor_horas"] = m_div.group(1)
            print(f"[POSTP] divisor_horas detectado: {cleaned['divisor_horas']}", flush=True)
        else:
            jornada = cleaned.get("jornada_contratual") or ""
            m_sem = re.search(r"(\d+)\s*h(?:oras?)?\s*semanais?", jornada, re.IGNORECASE)
            if m_sem:
                h_sem = int(m_sem.group(1))
                divisores = {30: "150", 35: "175", 36: "180", 40: "200", 44: "220"}
                if h_sem in divisores:
                    cleaned["divisor_horas"] = divisores[h_sem]
                    print(f"[POSTP] divisor_horas inferido da jornada ({h_sem}h/sem): {cleaned['divisor_horas']}", flush=True)
                else:
                    import math
                    divisor_calc = str(math.ceil((h_sem / 6) * 30))
                    cleaned["divisor_horas"] = divisor_calc
                    print(f"[POSTP] divisor_horas calculado matematicamente ({h_sem}h/sem -> {divisor_calc})", flush=True)

    # 5. evolucao_salarial — valor padrão se não extraído
    if not cleaned.get("evolucao_salarial"):
        salario = cleaned.get("salario_base")
        if salario:
            cleaned["evolucao_salarial"] = f"Salário fixo reconhecido: {salario}"

    return cleaned


# Campos obrigatórios mínimos para considerar extração completa
_CAMPOS_OBRIGATORIOS = [
    "numero_processo",
    "reclamante",
    "reclamada",
    "data_sentenca",
    "salario_base",
]
# Mínimo de verbas esperadas para uma sentença procedente típica
_MIN_VERBAS = 3


def _qualidade_ok(dados: dict, doc_type: str = "sentenca") -> tuple[bool, str]:
    """
    Verifica se a extração tem qualidade mínima para ser cacheada.
    Retorna (ok, motivo).

    - `liquidacao`: não exige salario_base nem mínimo de 3 verbas; exige partes ou
      número do processo e pelo menos uma verba.
    - Demais tipos: critérios rígidos (sentença / acórdão / dossiê completo).
    """
    dt = (doc_type or "sentenca").strip().lower()

    if dt == "peticao_inicial":
        ped = dados.get("verbas_pedidas") or []
        alt = dados.get("verbas_deferidas") or []
        if len(ped) + len(alt) < 1:
            return False, "peticao_inicial: nenhum pedido de verba identificado"
        return True, "ok"

    if dt == "contestacao":
        teses = dados.get("teses_defesa") or []
        if len(teses) < 1:
            return False, "contestacao: nenhuma tese de defesa identificada"
        return True, "ok"

    if dt == "liquidacao":
        np = dados.get("numero_processo")
        has_proc = bool(np and str(np).strip())
        rec_a = dados.get("reclamante")
        rec_b = dados.get("reclamada")
        has_partes = (bool(rec_a and str(rec_a).strip())) or (
            bool(rec_b and str(rec_b).strip())
        )
        if not has_proc and not has_partes:
            return (
                False,
                "liquidacao: falta identificação mínima (número do processo ou partes)",
            )
        n_verbas = len(dados.get("verbas_deferidas") or [])
        if n_verbas < 1:
            return False, "liquidacao: verbas_deferidas vazias"
        return True, "ok"

    faltando = [c for c in _CAMPOS_OBRIGATORIOS if not dados.get(c)]
    if faltando:
        return False, f"campos obrigatórios ausentes: {faltando}"

    n_verbas = len(dados.get("verbas_deferidas") or [])
    if n_verbas < _MIN_VERBAS:
        return False, f"apenas {n_verbas} verbas extraídas (mínimo: {_MIN_VERBAS})"

    return True, "ok"


def _vazio_ia(val) -> bool:
    """
    Retorna True se o valor retornado pela IA deve ser tratado como vazio.
    Usa .lower() para cobrir variações de capitalização (ex.: 'Não Informado').
    """
    if val is None:
        return True
    if isinstance(val, str) and val.strip().lower() in ("", "não informado", "nao informado"):
        return True
    return False


def _extrair_dados_regex_cabecalho(texto: str, label: str = "") -> dict:
    """
    Extrai data_ajuizamento, valor_causa e data_sentenca do cabeçalho/fim PJe via regex.
    Comum aos dois fluxos (PDF e dossiê).

    Args:
        texto: texto completo do documento.
        label: sufixo opcional para o print de log (ex: " (dossiê)").
    """
    cabecalho = texto[:3000]
    final_doc = texto[-2500:] if len(texto) > 2500 else texto
    dados_regex: dict = {}

    m_autuacao = re.search(r"Data da Autuação:\s*(\d{2}/\d{2}/\d{4})", cabecalho)
    if m_autuacao:
        dados_regex["data_ajuizamento"] = m_autuacao.group(1)

    m_valor = re.search(r"Valor da causa:\s*(R\\?\$?\s*[\d\.,]+)", cabecalho)
    if m_valor:
        valor = m_valor.group(1).strip().replace("\\", "").strip()
        if valor:
            dados_regex["valor_causa"] = valor

    m_assinado = re.search(
        r"Assinado\s+eletronicamente\s+em\s+(\d{2}/\d{2}/\d{4})", cabecalho, re.IGNORECASE
    )
    if m_assinado:
        dados_regex["data_sentenca"] = m_assinado.group(1)

    if "data_sentenca" not in dados_regex:
        m_julg = re.search(
            r"Data\s+do\s+Julgamento:\s*(\d{2}/\d{2}/\d{4})", final_doc, re.IGNORECASE
        )
        if m_julg:
            dados_regex["data_sentenca"] = m_julg.group(1)

    if "data_sentenca" not in dados_regex:
        m_pub = re.search(
            r"Publicado\s+em\s+(\d{2}/\d{2}/\d{4})", final_doc, re.IGNORECASE
        )
        if m_pub:
            dados_regex["data_sentenca"] = m_pub.group(1)

    if dados_regex:
        print(f"[PROCESSOR] Regex cabeçalho/fim PJe{label}: {list(dados_regex.keys())}", flush=True)

    return dados_regex


def _coletar_fontes_extracao(
    dados_finais: dict,
    texto: str,
    dados_regex: Optional[dict] = None,
) -> list[dict]:
    """Monta trilha de proveniência dos campos exibidos ao usuário."""
    try:
        out: list[dict] = []
        dados_regex = dados_regex or {}
        chunks = build_page_chunks(texto or "")
        max_page = max(chunks.keys()) if chunks else None

        def _append(
            campo: str,
            valor: Any,
            *,
            pagina: Optional[int] = None,
            trecho: Optional[str] = None,
            origem: str = "ia",
            confianca: Optional[float] = None,
        ) -> None:
            if valor is None:
                return
            valor_txt = str(valor).strip()
            if not valor_txt:
                return
            out.append(
                {
                    "campo": campo,
                    "valor_resumo": valor_txt[:240],
                    "pagina_origem": pagina,
                    "trecho": (str(trecho).strip()[:300] if trecho else None),
                    "origem": origem,
                    "confianca": confianca,
                }
            )

        for campo in (
            "numero_processo",
            "reclamante",
            "reclamada",
            "data_sentenca",
            "data_ajuizamento",
            "data_admissao",
            "data_demissao",
            "salario_base",
            "valor_causa",
        ):
            origem = "regex" if campo in dados_regex else "ia"
            pagina = None
            if campo in dados_regex:
                if campo in ("data_ajuizamento", "valor_causa"):
                    pagina = 1
                elif campo == "data_sentenca" and max_page is not None:
                    pagina = max_page
            _append(campo, dados_finais.get(campo), pagina=pagina, origem=origem)

        for i, verba in enumerate(dados_finais.get("verbas_deferidas") or []):
            if not isinstance(verba, dict):
                continue
            pagina = verba.get("pagina_origem")
            trecho = verba.get("trecho_fundamentacao")
            _append(f"verbas_deferidas[{i}].nome", verba.get("nome"), pagina=pagina, trecho=trecho)
            _append(f"verbas_deferidas[{i}].periodo", verba.get("periodo"), pagina=pagina, trecho=trecho)
            _append(
                f"verbas_deferidas[{i}].observacoes",
                verba.get("observacoes"),
                pagina=pagina,
                trecho=trecho,
            )

        for i, tese in enumerate(dados_finais.get("teses_defesa") or []):
            if not isinstance(tese, dict):
                continue
            pagina = tese.get("pagina_origem")
            trecho = tese.get("trecho_fundamentacao")
            _append(f"teses_defesa[{i}].verba_alvo", tese.get("verba_alvo"), pagina=pagina, trecho=trecho)
            _append(
                f"teses_defesa[{i}].tese_principal",
                tese.get("tese_principal"),
                pagina=pagina,
                trecho=trecho,
            )

        for i, item in enumerate(dados_finais.get("quadro_comparativo") or []):
            if not isinstance(item, dict):
                continue
            _append(f"quadro_comparativo[{i}].verba_alvo", item.get("verba_alvo"))
            _append(f"quadro_comparativo[{i}].resumo_pedido", item.get("resumo_pedido"))
            _append(f"quadro_comparativo[{i}].resumo_defesa", item.get("resumo_defesa"))
            _append(f"quadro_comparativo[{i}].resumo_decisao", item.get("resumo_decisao"))
            _append(f"quadro_comparativo[{i}].status_final", item.get("status_final"))

        # Alertas jurídicos: tentativa de apontar fonte principal do alerta.
        verbas_idx: list[tuple[str, Optional[int], Optional[str]]] = []
        for verba in (dados_finais.get("verbas_deferidas") or []):
            if not isinstance(verba, dict):
                continue
            nome = str(verba.get("nome") or "").strip().lower()
            if not nome:
                continue
            verbas_idx.append(
                (nome, verba.get("pagina_origem"), verba.get("trecho_fundamentacao"))
            )

        for idx, alerta in enumerate(dados_finais.get("alertas_juridicos") or []):
            if not isinstance(alerta, str):
                continue
            alerta_l = alerta.lower()
            pagina: Optional[int] = None
            trecho: Optional[str] = None
            confianca: Optional[float] = None

            # 1) Tenta casar verba por nome explícito (ou entre aspas)
            quoted = re.findall(r"[\"']([^\"']+)[\"']", alerta)
            candidatos = [q.strip().lower() for q in quoted if str(q).strip()]
            for nome_verba, pag_verba, tre_verba in verbas_idx:
                if nome_verba in alerta_l or any(c in nome_verba for c in candidatos):
                    pagina = pag_verba
                    trecho = tre_verba
                    confianca = 0.95 if pagina is not None else 0.8
                    break

            # 2) Regras de identificação por tema do alerta
            if pagina is None and ("cnj" in alerta_l or "processo" in alerta_l):
                pagina = 1
                confianca = 0.7
            if pagina is None and "ajuiz" in alerta_l:
                pagina = 1
                confianca = 0.6
            if pagina is None and ("senten" in alerta_l or "julgamento" in alerta_l):
                pagina = max_page
                confianca = 0.6 if max_page is not None else 0.4
            if pagina is None and ("admiss" in alerta_l or "demiss" in alerta_l):
                pagina = 1
                confianca = 0.5

            _append(
                f"alertas_juridicos[{idx}]",
                alerta,
                pagina=pagina,
                trecho=trecho,
                origem="regra",
                confianca=confianca,
            )

        return out
    except Exception:
        return []


def _emit_ws_partial(job_id: str, payload: dict, message: str = "") -> None:
    """Envia fatia JSON para /ws/{job_id} quando há conexão (lazy import evita ciclo)."""
    jid = (job_id or "").strip()
    if not jid or not payload:
        return
    try:
        from api.routers import extractor as _ext

        _ext.push_ws_partial(jid, payload, message)
    except Exception:
        pass


def _partial_payload_pos_ia(d: dict, doc_type: str) -> dict:
    """Campos principais do ProcessoTrabalhista já após IA + validação (antes ou após ancoragem)."""
    keys_scalar = (
        "numero_processo",
        "reclamante",
        "reclamada",
        "vara_trabalho",
        "data_sentenca",
        "data_ajuizamento",
        "data_admissao",
        "data_demissao",
        "salario_base",
        "motivo_rescisao",
        "aviso_previo_dias",
        "valor_causa",
    )
    out: dict = {"_meta_doc_type": doc_type}
    for k in keys_scalar:
        v = d.get(k)
        if v is not None and str(v).strip() != "":
            out[k] = v
    for k in ("verbas_deferidas", "verbas_pedidas", "quadro_comparativo", "teses_defesa"):
        if k in d and d.get(k) is not None:
            out[k] = d[k]
    if d.get("dano_moral") is not None:
        out["dano_moral"] = d["dano_moral"]
    return out


def _partial_payload_pos_motor(dados_finais: dict) -> dict:
    return {
        "alertas_juridicos": dados_finais.get("alertas_juridicos") or [],
        "memorial_juridico": dados_finais.get("memorial_juridico"),
        "regras_aplicadas": dados_finais.get("regras_aplicadas") or [],
        "explicacoes": dados_finais.get("explicacoes") or [],
        "fontes_extracao": dados_finais.get("fontes_extracao") or [],
    }


def _partial_payload_pos_parecer(dados_finais: dict) -> dict:
    keys = (
        "parecer_texto",
        "parecer_parcelas_ia",
        "parecer_intro_parcelas",
        "parecer_parcelas_apuradas",
        "parecer_model_used",
        "parecer_criterios_inss",
        "parecer_criterios_irrf",
    )
    out: dict = {}
    for k in keys:
        if k not in dados_finais:
            continue
        v = dados_finais.get(k)
        if v is None:
            continue
        if v == "" or v == []:
            continue
        out[k] = v
    return out


def _executar_pipeline_pos_ia(
    texto: str,
    dados_limpos: dict,
    avisos_dedup: list,
    ai_result: dict,
    doc_type: str,
    user_id: str,
    job_id: str,
    cache_repo,
    extraction_repo,
    cache_hash: str,
    label: str = "",
) -> dict:
    """
    Passos 4b–10 do pipeline: compartilhados por process_lawsuit_pdf e
    process_lawsuit_dossie. Não deve ser chamado diretamente por código externo.

    Args:
        texto:           Texto completo do documento (usado para regex cabeçalho).
        dados_limpos:    Saída de _validate_result (verbas já deduplicadas).
        avisos_dedup:    Avisos de deduplicação de verbas (pode ser lista vazia).
        ai_result:       Dict {"data", "model_used", "error"} retornado pelo ai_client.
        doc_type:        Tipo de documento detectado ("sentenca", "completo", etc.).
        user_id:         ID do usuário/tenant.
        job_id:          ID do job (vindo do main.py; fallback para doc_id do banco).
        cache_repo:      Repositório de cache instanciado pelo caller.
        extraction_repo: Repositório de extrações instanciado pelo caller.
        cache_hash:      Hash SHA-256 do PDF ou do dossiê para salvar no cache.
        label:           Sufixo de log para distinguir PDF de dossiê (ex: " (dossiê)").

    Returns:
        Dict de resposta final no mesmo formato de process_lawsuit_pdf/dossie.
    """
    # 4b. Regex cabeçalho — mescla campos ausentes/vazios vindos da IA
    dados_regex = _extrair_dados_regex_cabecalho(texto, label=label)

    # 7. Validação Pydantic
    try:
        processo = ProcessoTrabalhista(**dados_limpos)
        dados_finais = processo.model_dump()
    except Exception as e:
        return {"status": "erro", "msg": f"Dados inválidos da IA: {e}"}

    # 7b. Mesclagem camada Regex: preenche campos se IA retornou vazio
    for key in ("data_ajuizamento", "valor_causa", "data_sentenca"):
        if key in dados_regex and _vazio_ia(dados_finais.get(key)):
            dados_finais[key] = dados_regex[key]
            print(f"[PROCESSOR] Mescla Regex → {key}", flush=True)

    # 7c. Fonte por página: cruza trecho_fundamentacao com marcadores --- PÁGINA N ---
    anchor_verbas_to_pages(texto, dados_finais.get("verbas_deferidas") or [])
    _emit_ws_partial(
        job_id,
        _partial_payload_pos_ia(dados_finais, doc_type),
        "[PIPELINE] Cabeçalho e verbas ancoradas (pré-motor jurídico).",
    )
    # 8a. LegalRuleEngine estático (via interface legal_validator):
    # evita execução duplicada do engine no mesmo request.
    resultado_engine = validar_dados_completo(dados_finais)
    alertas_engine = resultado_engine["alertas"]
    regras_aplicadas = resultado_engine["regras_aplicadas"]
    memorial_juridico = resultado_engine["memorial_juridico"]

    # 8b. Self-Healing: injeta regras dinâmicas ATIVAS (Knowledge Base)
    try:
        regras_dinamicas = carregar_regras_ativas()
        if regras_dinamicas:
            engine_dyn = LegalRuleEngine(regras_dinamicas)
            res_dyn = engine_dyn.executar(dados_finais)
            alertas_engine   = alertas_engine + res_dyn.get("alertas", [])
            regras_aplicadas = regras_aplicadas + res_dyn.get("regras_aplicadas", [])
    except Exception as _e_dyn:
        print(f"[PROCESSOR] Aviso: erro nas regras dinâmicas (não crítico): {_e_dyn}", flush=True)

    # 8c. Explanation Engine: texto jurídico por verba (sem LLM)
    explicacoes = gerar_explicacoes(
        verbas=dados_finais.get("verbas_deferidas", []),
        memorial_juridico=memorial_juridico,
    )

    # Mescla todos os alertas (sem duplicatas textuais)
    dados_finais["alertas_juridicos"] = _dedup_alertas(alertas_engine, avisos_dedup)
    dados_finais["regras_aplicadas"]  = regras_aplicadas
    dados_finais["memorial_juridico"] = memorial_juridico
    dados_finais["explicacoes"]       = explicacoes
    dados_finais["fontes_extracao"] = _coletar_fontes_extracao(
        dados_finais, texto, dados_regex
    )
    _emit_ws_partial(
        job_id,
        _partial_payload_pos_motor(dados_finais),
        "[PIPELINE] Alertas, memorial, regras e rastreabilidade (fontes) consolidados.",
    )

    # Parecer técnico — I. PARCELAS APURADAS
    parecer_secao_i = gerar_parecer_parcelas_apuradas(
        dados_finais.get("verbas_deferidas", []),
        dados_finais,
    )
    dados_finais["parecer_intro_parcelas"]    = parecer_secao_i.get("intro", "")
    dados_finais["parecer_parcelas_apuradas"] = parecer_secao_i.get("itens", [])

    # Textos padrão da seção II. CRITÉRIOS UTILIZADOS
    criterios_padrao = obter_textos_padrao_criterios_parecer()
    dados_finais["parecer_criterios_inss"] = criterios_padrao.get("inss", "")
    dados_finais["parecer_criterios_irrf"] = criterios_padrao.get("irrf", "")

    # Parecer técnico completo — template fixo + slot gerado pela IA no Padrão Ouro
    try:
        parecer_completo = gerar_parecer_tecnico_completo(
            dados_finais,
            dados_finais.get("verbas_deferidas", []),
        )
        dados_finais["parecer_texto"]       = parecer_completo.get("texto", "")
        dados_finais["parecer_parcelas_ia"] = parecer_completo.get("parcelas", "")
        dados_finais["parecer_model_used"]  = parecer_completo.get("model_used")
        if parecer_completo.get("error"):
            print(
                f"[PARECER] Aviso: IA retornou erro ao gerar parcelas — {parecer_completo['error']}",
                flush=True,
            )
    except Exception as _e_parecer:
        dados_finais["parecer_texto"]       = ""
        dados_finais["parecer_parcelas_ia"] = ""
        dados_finais["parecer_model_used"]  = None
        print(f"[PARECER] Erro ao gerar parecer completo (não crítico): {_e_parecer}", flush=True)

    _pp = _partial_payload_pos_parecer(dados_finais)
    if _pp:
        _emit_ws_partial(
            job_id,
            _pp,
            "[PIPELINE] Parecer técnico e critérios (seções geradas).",
        )

    # 8d. Shadow Mode: regras KB status=shadow (observação; persiste em shadow_logs)
    try:
        dados_finais["shadow_logs"] = executar_shadow_pipeline(dados_finais)
    except Exception as _e_shadow:
        print(f"[SHADOW] Aviso: erro no shadow pipeline (não crítico): {_e_shadow}", flush=True)
        dados_finais["shadow_logs"] = []

    # 9. Só cacheia e desconta crédito se qualidade mínima atingida
    ok, motivo = _qualidade_ok(dados_finais, doc_type)
    if ok:
        dados_finais["_meta_doc_type"] = doc_type
        cache_repo.save_cache(cache_hash, dados_finais)
        doc_id = extraction_repo.save_extraction(
            user_id=user_id,
            data=dados_finais,
            doc_type=doc_type,
            model_used=ai_result["model_used"],
        )
        deduct_credit(user_id)
        print(f"[PROCESSOR] Resultado salvo{label} (qualidade ok)", flush=True)
    else:
        doc_id = extraction_repo.save_extraction(
            user_id=user_id,
            data=dados_finais,
            doc_type=doc_type,
            model_used=ai_result["model_used"],
        )
        print(
            f"[PROCESSOR] Qualidade insuficiente{label} — NÃO cacheado: {motivo}",
            flush=True,
        )

    # 10. Memória de cálculo
    gerar_memoria(
        job_id=job_id or str(doc_id),
        dados=dados_finais,
        model_used=ai_result["model_used"],
        doc_type=doc_type,
        avisos_dedup=avisos_dedup,
        explicacoes=explicacoes,
    )

    raiox = enriquecer_para_raiox(dados_finais)

    return {
        "status":            "sucesso",
        "source":            "ai",
        "doc_type":          doc_type,
        "model_used":        ai_result["model_used"],
        "doc_id":            doc_id,
        "data":              dados_finais,
        "raiox":             raiox,
        "alertas_juridicos": dados_finais["alertas_juridicos"],
        "regras_aplicadas":  regras_aplicadas,
        "memorial_juridico": memorial_juridico,
        "explicacoes":       explicacoes,
        "qualidade_ok":      ok,
        "qualidade_motivo":  motivo if not ok else None,
    }


def _verba_dict_from_pedido_item(item: Any) -> Optional[dict]:
    """Normaliza um item de verbas_pedidas (string ou dict) para dict compatível com VerbaDeferida."""
    if isinstance(item, str):
        name = item.strip()
        if not name:
            return None
        return {"nome": name, "status_final": "pedido", "reflexos": []}
    if isinstance(item, dict):
        nome = (
            str(item.get("nome") or item.get("verba") or item.get("pedido") or "")
        ).strip()
        if not nome:
            return None
        out: dict = {
            "nome": nome,
            "status_final": str(item.get("status_final") or "pedido").strip(),
            "reflexos": list(item.get("reflexos") or []),
        }
        for opt in ("observacoes", "periodo", "percentual", "trecho_fundamentacao"):
            if item.get(opt) is not None:
                out[opt] = item.get(opt)
        if item.get("pagina_origem") is not None:
            try:
                out["pagina_origem"] = int(item["pagina_origem"])
            except (TypeError, ValueError):
                pass
        return out
    return None


def _montar_base_peticao_inicial(raw: dict) -> dict:
    pedidas_raw = raw.get("verbas_pedidas") or []
    verbas: list = []
    for it in pedidas_raw:
        vd = _verba_dict_from_pedido_item(it)
        if vd:
            verbas.append(vd)
    vc = raw.get("valor_causa")
    vc_str: Optional[str] = None
    if vc is not None and str(vc).strip():
        vc_str = str(vc).strip()
    return {
        "verbas_deferidas": verbas,
        "verbas_pedidas": verbas,
        "valor_causa": vc_str,
    }


def _tese_item_para_dict(item: Any) -> Optional[dict]:
    """Normaliza item de teses_defesa (IA/Lab) para dict compatível com TeseDefesa."""
    if not isinstance(item, dict):
        return None
    alvo = str(item.get("verba_alvo") or item.get("verba") or "").strip()
    tese = str(item.get("tese_principal") or item.get("argumento") or "").strip()
    if not alvo and not tese:
        return None
    trecho_raw = item.get("trecho_fundamentacao")
    ts: Optional[str] = None
    if trecho_raw is not None and str(trecho_raw).strip():
        ts = str(trecho_raw).strip()
        if len(ts) > 200:
            ts = ts[:200]
    po = item.get("pagina_origem")
    pagina: Optional[int] = None
    if po is not None:
        try:
            pagina = int(po)
        except (TypeError, ValueError):
            pagina = None
    return {
        "verba_alvo": alvo or "Pedido não especificado",
        "tese_principal": tese or "Tese não especificada",
        "trecho_fundamentacao": ts,
        "pagina_origem": pagina,
        "incontroversa": bool(item.get("incontroversa")),
    }


def _montar_base_contestacao(raw: dict) -> dict:
    def _opt_str(val: Any) -> Optional[str]:
        if val is None:
            return None
        s = str(val).strip()
        return s or None

    teses_in = raw.get("teses_defesa") or []
    teses: list = []
    for it in teses_in:
        td = _tese_item_para_dict(it)
        if td:
            teses.append(td)
    return {
        "numero_processo": _opt_str(raw.get("numero_processo")),
        "reclamante": _opt_str(raw.get("reclamante")),
        "reclamada": _opt_str(raw.get("reclamada")),
        "valor_causa": _opt_str(raw.get("valor_causa")),
        "verbas_deferidas": [],
        "verbas_pedidas": None,
        "teses_defesa": teses,
    }


def _pipeline_contestacao(
    user_id: str,
    file_bytes: bytes,
    job_id: str,
    filename: str,
    cache_repo,
    extraction_repo,
) -> dict:
    """Fluxo dedicado: contestação (sem sentence_finder + motor de sentença)."""
    from services.learning_engine import (
        _extrair_contestacao,
        _extrair_texto_arquivo_com_marcadores_pagina,
    )

    pdf_hash = _hash_pdf(file_bytes)
    storage_key = pdf_cache_storage_key(pdf_hash, "contestacao")
    cached = cache_repo.get_cache(storage_key)
    if cached:
        ok_c, _mot = _qualidade_ok(cached, "contestacao")
        if ok_c:
            print("[PROCESSOR] Cache hit (contestacao)", flush=True)
            return {
                "status": "sucesso",
                "source": "cache",
                "doc_type": "contestacao",
                "data": cached,
            }
        print(f"[PROCESSOR] Cache contestação descartado: {_mot}", flush=True)

    raw = _extrair_contestacao(file_bytes, filename)
    if raw.get("erro"):
        return {"status": "erro", "msg": str(raw["erro"])}

    base = _montar_base_contestacao(raw)
    try:
        processo = ProcessoTrabalhista(**base)
        dados_finais = processo.model_dump()
    except Exception as e:
        return {"status": "erro", "msg": f"Dados inválidos da contestação: {e}"}

    texto_marcado = _extrair_texto_arquivo_com_marcadores_pagina(
        file_bytes, filename
    )
    if texto_marcado:
        anchor_verbas_to_pages(texto_marcado, dados_finais.get("teses_defesa") or [])
    dados_finais["fontes_extracao"] = _coletar_fontes_extracao(
        dados_finais, texto_marcado
    )

    memorial = gerar_memorial_defesa(
        dados_finais.get("teses_defesa"),
        reclamada=str(dados_finais.get("reclamada") or ""),
    )
    dados_finais["memorial_juridico"] = memorial
    dados_finais["explicacoes"] = []
    dados_finais["alertas_juridicos"] = []
    dados_finais["regras_aplicadas"] = []
    dados_finais["parecer_texto"] = ""
    dados_finais["parecer_parcelas_ia"] = ""
    dados_finais["parecer_intro_parcelas"] = ""
    dados_finais["parecer_parcelas_apuradas"] = []
    dados_finais["parecer_model_used"] = None
    dados_finais["parecer_criterios_inss"] = ""
    dados_finais["parecer_criterios_irrf"] = ""
    dados_finais["_meta_doc_type"] = "contestacao"

    try:
        dados_finais["shadow_logs"] = executar_shadow_pipeline(dados_finais)
    except Exception as _e_sh:
        print(f"[SHADOW] Aviso contestação: {_e_sh}", flush=True)
        dados_finais["shadow_logs"] = []

    pl_c = _partial_payload_pos_ia(dados_finais, "contestacao")
    pl_c["memorial_juridico"] = memorial
    if dados_finais.get("shadow_logs"):
        pl_c["shadow_logs"] = dados_finais["shadow_logs"]
    _emit_ws_partial(job_id, pl_c, "[CONTESTAÇÃO] Teses, memorial de defesa e trilha consolidados.")

    model_used = raw.get("model_used") or "gemini"
    ok, motivo = _qualidade_ok(dados_finais, "contestacao")

    if ok:
        cache_repo.save_cache(storage_key, dados_finais)
        print("[PROCESSOR] Resultado contestação salvo em cache", flush=True)

    doc_id = extraction_repo.save_extraction(
        user_id=user_id,
        data=dados_finais,
        doc_type="contestacao",
        model_used=str(model_used),
    )
    if ok:
        deduct_credit(user_id)

    try:
        gerar_memoria(
            job_id=job_id or str(doc_id),
            dados=dados_finais,
            model_used=str(model_used),
            doc_type="contestacao",
            avisos_dedup=[],
            explicacoes=[],
        )
    except Exception as _e_m:
        print(f"[MEMORIA] Aviso contestação: {_e_m}", flush=True)

    raiox = enriquecer_para_raiox(dados_finais)

    return {
        "status": "sucesso",
        "source": "ai",
        "doc_type": "contestacao",
        "model_used": model_used,
        "doc_id": doc_id,
        "data": dados_finais,
        "raiox": raiox,
        "alertas_juridicos": dados_finais["alertas_juridicos"],
        "regras_aplicadas": dados_finais["regras_aplicadas"],
        "memorial_juridico": memorial,
        "explicacoes": [],
        "qualidade_ok": ok,
        "qualidade_motivo": motivo if not ok else None,
    }


def _pipeline_peticao_inicial(
    user_id: str,
    file_bytes: bytes,
    job_id: str,
    filename: str,
    cache_repo,
    extraction_repo,
) -> dict:
    """Fluxo dedicado: petição inicial (sem sentence_finder + motor de sentença)."""
    pdf_hash = _hash_pdf(file_bytes)
    storage_key = pdf_cache_storage_key(pdf_hash, "peticao_inicial")
    cached = cache_repo.get_cache(storage_key)
    if cached:
        ok_c, _mot = _qualidade_ok(cached, "peticao_inicial")
        if ok_c:
            print("[PROCESSOR] Cache hit (peticao_inicial)", flush=True)
            return {
                "status": "sucesso",
                "source": "cache",
                "doc_type": "peticao_inicial",
                "data": cached,
            }
        print(f"[PROCESSOR] Cache petição descartado: {_mot}", flush=True)

    from services.learning_engine import _extrair_peticao_inicial

    raw = _extrair_peticao_inicial(file_bytes, filename)
    if raw.get("erro"):
        return {"status": "erro", "msg": str(raw["erro"])}

    base = _montar_base_peticao_inicial(raw)
    try:
        processo = ProcessoTrabalhista(**base)
        dados_finais = processo.model_dump()
    except Exception as e:
        return {"status": "erro", "msg": f"Dados inválidos da petição: {e}"}

    memorial = gerar_memorial_pedidos(
        dados_finais.get("verbas_pedidas"),
        causa_pedir=str(raw.get("causa_pedir") or ""),
        periodo_reivindicado=str(raw.get("periodo_reivindicado") or ""),
    )
    dados_finais["memorial_juridico"] = memorial
    dados_finais["explicacoes"] = []
    dados_finais["alertas_juridicos"] = []
    dados_finais["regras_aplicadas"] = []
    dados_finais["parecer_texto"] = ""
    dados_finais["parecer_parcelas_ia"] = ""
    dados_finais["parecer_intro_parcelas"] = ""
    dados_finais["parecer_parcelas_apuradas"] = []
    dados_finais["parecer_model_used"] = None
    dados_finais["parecer_criterios_inss"] = ""
    dados_finais["parecer_criterios_irrf"] = ""
    dados_finais["_meta_doc_type"] = "peticao_inicial"
    dados_finais["fontes_extracao"] = _coletar_fontes_extracao(dados_finais, texto="")

    try:
        dados_finais["shadow_logs"] = executar_shadow_pipeline(dados_finais)
    except Exception as _e_sh:
        print(f"[SHADOW] Aviso petição: {_e_sh}", flush=True)
        dados_finais["shadow_logs"] = []

    pl_p = _partial_payload_pos_ia(dados_finais, "peticao_inicial")
    pl_p["memorial_juridico"] = memorial
    if dados_finais.get("shadow_logs"):
        pl_p["shadow_logs"] = dados_finais["shadow_logs"]
    _emit_ws_partial(job_id, pl_p, "[PETIÇÃO INICIAL] Pedidos, memorial e trilha consolidados.")

    model_used = raw.get("model_used") or "gemini"
    ok, motivo = _qualidade_ok(dados_finais, "peticao_inicial")

    if ok:
        cache_repo.save_cache(storage_key, dados_finais)
        print("[PROCESSOR] Resultado petição salvo em cache", flush=True)

    doc_id = extraction_repo.save_extraction(
        user_id=user_id,
        data=dados_finais,
        doc_type="peticao_inicial",
        model_used=str(model_used),
    )
    if ok:
        deduct_credit(user_id)

    try:
        gerar_memoria(
            job_id=job_id or str(doc_id),
            dados=dados_finais,
            model_used=str(model_used),
            doc_type="peticao_inicial",
            avisos_dedup=[],
            explicacoes=[],
        )
    except Exception as _e_m:
        print(f"[MEMORIA] Aviso petição: {_e_m}", flush=True)

    raiox = enriquecer_para_raiox(dados_finais)

    return {
        "status": "sucesso",
        "source": "ai",
        "doc_type": "peticao_inicial",
        "model_used": model_used,
        "doc_id": doc_id,
        "data": dados_finais,
        "raiox": raiox,
        "alertas_juridicos": dados_finais["alertas_juridicos"],
        "regras_aplicadas": dados_finais["regras_aplicadas"],
        "memorial_juridico": memorial,
        "explicacoes": [],
        "qualidade_ok": ok,
        "qualidade_motivo": motivo if not ok else None,
    }


def process_lawsuit_pdf(
    user_id: str,
    file_bytes: bytes,
    job_id: str = "",
    *,
    cache_context: str = "auto",
    filename: str = "documento.pdf",
) -> dict:
    """
    Pipeline completo de extração.

    Parâmetros
    ----------
    user_id    : identificador do usuário (freemium / créditos)
    file_bytes : bytes do PDF
    job_id     : identificador do job vindo do main.py — usado na memória de cálculo (M1).
                 Opcional: se vazio, usa doc_id como fallback após salvar no banco.
    cache_context : "auto" = chave de cache só pelo hash (legado). "peticao_inicial" /
                    "contestacao" = chave composta + fluxo dedicado (sem motor de sentença).
    filename   : nome do ficheiro (extensão relevante para petição inicial / contestação).
    """
    print(f"[PROCESSOR] user={user_id}", flush=True)

    # 0. Repositórios — Sprint 2: tenant_id == user_id (separação futura em RequestContext)
    cache_repo = get_cache_repo()
    extraction_repo = get_extraction_repo(tenant_id=user_id)

    # 1. Freemium — verifica créditos
    credits = get_user_credits(user_id)
    if credits <= 0:
        return {"status": "erro", "msg": "Saldo esgotado. Adquira mais créditos."}

    # Sprint 5: validação de quota por plano (PDFs)
    if quota_excedida(user_id, "pdfs"):
        return {
            "status": "erro",
            "msg": "Limite de PDFs do plano atingido. Faça upgrade para continuar.",
        }

    ctx_req = (cache_context or "auto").strip().lower()
    if ctx_req == "peticao_inicial":
        fn = (filename or "documento.pdf").lower()
        if not (fn.endswith(".pdf") or fn.endswith(".docx") or fn.endswith(".doc")):
            return {
                "status": "erro",
                "msg": "Petição inicial: envie PDF, DOC ou DOCX.",
            }
        return _pipeline_peticao_inicial(
            user_id, file_bytes, job_id, filename, cache_repo, extraction_repo
        )

    if ctx_req == "contestacao":
        fn_c = (filename or "documento.pdf").lower()
        if not (fn_c.endswith(".pdf") or fn_c.endswith(".docx") or fn_c.endswith(".doc")):
            return {
                "status": "erro",
                "msg": "Contestação: envie PDF, DOC ou DOCX.",
            }
        return _pipeline_contestacao(
            user_id, file_bytes, job_id, filename, cache_repo, extraction_repo
        )

    # 2. Cache — só usa se tiver qualidade mínima
    pdf_hash = _hash_pdf(file_bytes)
    storage_key = pdf_cache_storage_key(pdf_hash, cache_context)
    cached = cache_repo.get_cache(storage_key)
    if cached:
        cached_doc_type = cached.get("_meta_doc_type", "sentenca")
        ok, motivo = _qualidade_ok(cached, cached_doc_type)
        if ok:
            print("[PROCESSOR] Cache hit! (qualidade ok)", flush=True)
            return {
                "status": "sucesso",
                "source": "cache",
                "doc_type": cached_doc_type,
                "data": cached,
            }
        else:
            print(f"[PROCESSOR] Cache descartado — qualidade insuficiente: {motivo}", flush=True)

    # 3. Extração de texto
    texto, doc_type = extract_sentence_from_pdf(file_bytes)
    if not texto.strip():
        return {"status": "erro", "msg": "PDF sem texto legível"}

    print(f"[PROCESSOR] Tipo detectado: {doc_type} | Chars extraídos: {len(texto)}", flush=True)

    pre_fields = pre_extract(texto)

    _pipeline_dbg_meta = {
        "job_id": job_id or "",
        "user_id": user_id,
        "label": f"single_pdf:{doc_type}",
    }
    if settings.DEBUG_PIPELINE:
        from services.pipeline_debug import log_etapa, maybe_write_dump

        log_etapa("step3_sentence_finder", texto, meta=_pipeline_dbg_meta)
        maybe_write_dump("step3_sentence_finder", texto, meta=_pipeline_dbg_meta)

    # 4. Playbook
    _PLAYBOOK_MAP = {
        "sentenca":   "sentenca_ordinaria.md",
        "acordao":    "acordao.md",
        "liquidacao": "calculo_liquidacao.md",
        "embargos":   "embargos_declaracao.md",
        "despacho":   "despacho_execucao.md",
        "completo":   "sentenca_ordinaria.md",
    }
    skill_file = _PLAYBOOK_MAP.get(doc_type, "sentenca_ordinaria.md")
    playbook   = _load_skill(skill_file)

    if find_section_hybrid(texto, "dispositivo") < 0:
        print("[SKILL] Dispositivo não encontrado — carregando filtro_dispositivo.md", flush=True)
        playbook += "\n\n" + _load_skill("filtro_dispositivo.md")

    # 5. IA
    ai_result = extract_data_with_gemini(
        texto,
        playbook=playbook,
        pre_fields=pre_fields,
        pipeline_debug_meta=_pipeline_dbg_meta,
    )
    if ai_result["error"]:
        return {"status": "erro", "msg": ai_result["error"]}

    # 6. Validação pós-IA
    dados_limpos = _validate_result(ai_result["data"])
    dados_limpos = _aplicar_pre_high(dados_limpos, pre_fields.get("high"))

    # 6b. Deduplicação de verbas
    if dados_limpos.get("verbas_deferidas"):
        verbas_dedup, avisos_dedup = deduplicar_verbas(dados_limpos["verbas_deferidas"])
        dados_limpos["verbas_deferidas"] = verbas_dedup
        if avisos_dedup:
            print(f"[DEDUP] {len(avisos_dedup)} duplicata(s) removida(s)", flush=True)
    else:
        avisos_dedup = []

    _emit_ws_partial(
        job_id,
        _partial_payload_pos_ia(dados_limpos, doc_type),
        "[PIPELINE] Resposta da IA validada e verbas deduplicadas.",
    )

    # 4b–10. Pipeline pós-IA (compartilhado)
    return _executar_pipeline_pos_ia(
        texto=texto,
        dados_limpos=dados_limpos,
        avisos_dedup=avisos_dedup,
        ai_result=ai_result,
        doc_type=doc_type,
        user_id=user_id,
        job_id=job_id,
        cache_repo=cache_repo,
        extraction_repo=extraction_repo,
        cache_hash=storage_key,
        label="",
    )


def process_lawsuit_dossie(user_id: str, files_list: list, job_id: str = "") -> dict:
    """
    Pipeline de extração para dossiê processual (múltiplos arquivos).
    files_list: [(filename, bytes), ...] — PDF, DOC, DOCX, PJC, XML.
    Extrai texto de cada arquivo, concatena com separadores e envia o super-contexto à IA.
    Cache por hash composto de todos os arquivos.
    """
    print(f"[PROCESSOR] Dossiê: user={user_id} | {len(files_list)} arquivo(s)", flush=True)

    # 0. Repositórios — Sprint 2: tenant_id == user_id
    cache_repo = get_cache_repo()
    extraction_repo = get_extraction_repo(tenant_id=user_id)

    # 1. Freemium
    credits = get_user_credits(user_id)
    if credits <= 0:
        return {"status": "erro", "msg": "Saldo esgotado. Adquira mais créditos."}

    # 2. Cache
    dossie_hash = _hash_dossie(files_list)
    cached      = cache_repo.get_cache(dossie_hash)
    if cached:
        cached_doc_type = cached.get("_meta_doc_type", "completo")
        ok, motivo = _qualidade_ok(cached, cached_doc_type)
        if ok:
            print("[PROCESSOR] Cache hit dossiê! (qualidade ok)", flush=True)
            return {
                "status":   "sucesso",
                "source":   "cache",
                "doc_type": cached_doc_type,
                "data":     cached,
            }
        else:
            print(f"[PROCESSOR] Cache dossiê descartado — {motivo}", flush=True)

    # 3. Super-contexto: extrai texto de cada arquivo e concatena
    partes = []
    for (nome, content) in files_list:
        txt = _extrair_texto_arquivo_dossie(nome, content)
        partes.append(f"\n--- INÍCIO DO DOCUMENTO: {nome} ---\n{txt}\n--- FIM DO DOCUMENTO ---\n")
    texto = "\n".join(partes)
    if not texto.strip():
        return {"status": "erro", "msg": "Nenhum texto legível nos arquivos do dossiê."}

    doc_type = "completo"
    print(f"[PROCESSOR] Dossiê: {len(texto)} chars (super-contexto)", flush=True)

    pre_fields = pre_extract(texto)

    _pipeline_dbg_meta_d = {
        "job_id": job_id or "",
        "user_id": user_id,
        "label": "dossie:completo",
    }
    if settings.DEBUG_PIPELINE:
        from services.pipeline_debug import log_etapa, maybe_write_dump

        log_etapa("step3_dossie_supercontexto", texto, meta=_pipeline_dbg_meta_d)
        maybe_write_dump("step3_dossie_supercontexto", texto, meta=_pipeline_dbg_meta_d)

    # 4. Playbook
    playbook = _load_skill("sentenca_ordinaria.md")
    if find_section_hybrid(texto, "dispositivo") < 0:
        playbook += "\n\n" + _load_skill("filtro_dispositivo.md")

    # 5. IA
    ai_result = extract_data_with_gemini(
        texto,
        playbook=playbook,
        pre_fields=pre_fields,
        pipeline_debug_meta=_pipeline_dbg_meta_d,
    )
    if ai_result["error"]:
        return {"status": "erro", "msg": ai_result["error"]}

    # 6. Validação pós-IA
    dados_limpos = _validate_result(ai_result["data"])
    dados_limpos = _aplicar_pre_high(dados_limpos, pre_fields.get("high"))

    # 6b. Deduplicação de verbas
    if dados_limpos.get("verbas_deferidas"):
        verbas_dedup, avisos_dedup = deduplicar_verbas(dados_limpos["verbas_deferidas"])
        dados_limpos["verbas_deferidas"] = verbas_dedup
        avisos_dedup = avisos_dedup or []
    else:
        avisos_dedup = []

    # 6c. Quadro comparativo (≥3 ficheiros): 2ª passagem IA — falha não bloqueia o fluxo
    dados_limpos["quadro_comparativo"] = []
    if len(files_list) >= 3:
        try:
            from services.ai_client import extrair_quadro_comparativo_dossie

            qr = extrair_quadro_comparativo_dossie(texto)
            qc = qr.get("quadro_comparativo") or []
            if isinstance(qc, list) and qc:
                dados_limpos["quadro_comparativo"] = qc
                print(
                    f"[DOSSIE] Quadro comparativo: {len(qc)} linha(s) "
                    f"(modelo={qr.get('model_used')})",
                    flush=True,
            )
        except Exception as _e_q:
            print(
                f"[DOSSIE] Quadro comparativo (não crítico): {_e_q}",
                flush=True,
            )

    _emit_ws_partial(
        job_id,
        _partial_payload_pos_ia(dados_limpos, doc_type),
        "[DOSSIÊ] Texto agregado; IA validada e quadro comparativo quando aplicável.",
    )

    # 4b–10. Pipeline pós-IA (compartilhado)
    return _executar_pipeline_pos_ia(
        texto=texto,
        dados_limpos=dados_limpos,
        avisos_dedup=avisos_dedup,
        ai_result=ai_result,
        doc_type=doc_type,
        user_id=user_id,
        job_id=job_id,
        cache_repo=cache_repo,
        extraction_repo=extraction_repo,
        cache_hash=dossie_hash,
        label=" (dossiê)",
    )