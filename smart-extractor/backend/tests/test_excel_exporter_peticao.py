"""Exportação Excel — ramo petição inicial (_meta_doc_type)."""

import os

from openpyxl import load_workbook

from services.excel_exporter import (
    exportar_excel,
    _is_peticao_inicial,
    _is_contestacao,
)


def test_is_contestacao():
    assert _is_contestacao({"_meta_doc_type": "contestacao"}) is True
    assert _is_contestacao({"_meta_doc_type": "CONTESTACAO"}) is True
    assert _is_contestacao({}) is False


def test_is_peticao_inicial():
    assert _is_peticao_inicial({"_meta_doc_type": "peticao_inicial"}) is True
    assert _is_peticao_inicial({"_meta_doc_type": "PETICAO_INICIAL"}) is True
    assert _is_peticao_inicial({}) is False
    assert _is_peticao_inicial({"_meta_doc_type": "sentenca"}) is False


def test_export_peticao_aba_pedidos_e_memorial():
    path = exportar_excel(
        {
            "_meta_doc_type": "peticao_inicial",
            "memorial_juridico": "Memorial sintético de teste.",
            "numero_processo": "0000000-00.0000.0.00.0000",
            "valor_causa": "R$ 10.000,00",
            "verbas_deferidas": [
                {
                    "nome": "Horas extras",
                    "status_final": "pedido",
                    "periodo": "2020",
                    "trecho_fundamentacao": "Pedido de HE",
                }
            ],
        },
        "job_test_pet",
    )
    try:
        wb = load_workbook(path, read_only=True)
        names = wb.sheetnames
        assert "Pedidos da inicial" in names
        assert "Resumo (Petição inicial)" in names
        wb.close()
    finally:
        try:
            os.unlink(path)
        except OSError:
            pass


def test_export_contestacao_aba_teses():
    path = exportar_excel(
        {
            "_meta_doc_type": "contestacao",
            "memorial_juridico": "Memorial de defesa teste.",
            "numero_processo": "0000000-00.0000.0.00.0000",
            "teses_defesa": [
                {
                    "verba_alvo": "HE",
                    "tese_principal": "Nega",
                    "trecho_fundamentacao": "x",
                    "incontroversa": False,
                    "pagina_origem": 2,
                }
            ],
        },
        "job_test_cont",
    )
    try:
        wb = load_workbook(path, read_only=True)
        names = wb.sheetnames
        assert "Teses de defesa" in names
        assert "Resumo (Contestação)" in names
        wb.close()
    finally:
        try:
            os.unlink(path)
        except OSError:
            pass


def test_export_com_quadro_comparativo_sheet():
    path = exportar_excel(
        {
            "numero_processo": "1",
            "verbas_deferidas": [
                {"nome": "FGTS", "status_final": "deferida", "reflexos": []},
            ],
            "quadro_comparativo": [
                {
                    "verba_alvo": "Horas extras",
                    "resumo_pedido": "Pedido",
                    "resumo_defesa": "Defesa",
                    "resumo_decisao": "Julgo",
                    "status_final": "Deferida",
                }
            ],
        },
        "job_quadro",
    )
    try:
        wb = load_workbook(path, read_only=True)
        assert "Quadro comparativo" in wb.sheetnames
        wb.close()
    finally:
        try:
            os.unlink(path)
        except OSError:
            pass


def test_export_sentenca_sheet_names_inalterados():
    path = exportar_excel(
        {
            "numero_processo": "1",
            "verbas_deferidas": [
                {"nome": "FGTS", "status_final": "deferida", "reflexos": []},
            ],
        },
        "job_test_sent",
    )
    try:
        wb = load_workbook(path, read_only=True)
        assert wb.sheetnames[0] == "Resumo e Parecer"
        assert "Verbas Deferidas" in wb.sheetnames
        wb.close()
    finally:
        try:
            os.unlink(path)
        except OSError:
            pass
